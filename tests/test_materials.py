import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from source_evidence import snapshot, read_json
from materials import build, audit, wrap_line
from init_project import initialize
from export_deposit import export

def fixture(base, large=False):
    from PIL import Image, ImageDraw
    base=Path(base);project=base/'project';project.mkdir()
    code='def calculate(values):\n    return sum(values)\n\nLABEL = "中文  spaces"\nDESCRIPTION = """first\n\n  significant spaces"""\nassert calculate([1, 2, 3]) == 6\n'
    if large: code+='\n'.join(f'assert calculate([{i}, {i+1}]) == {2*i+1}' for i in range(100))+'\n'
    code+='LONG = '+repr('测试 中文 原样  '+('ab  cd '*40))+'\n'
    (project/'main.py').write_text(code,encoding='utf-8')
    selection=base/'selection.json';selection.write_text('["main.py"]',encoding='utf-8')
    snap=base/'source.json';snapshot(project,selection,snap)
    image=Image.new('RGB',(900,360),'white');draw=ImageDraw.Draw(image)
    draw.rectangle((40,40,860,320),outline='black',width=3);draw.text((80,100),'TEST FIXTURE - NOT AN ACTUAL SOFTWARE SCREENSHOT',fill='black')
    image.save(base/'fixture.png')
    app={'name':'隔离验证软件','version':'V1.0','category':'应用软件','owners':[{'name':'测试主体'}],
         'completion_date':None,'development_mode':None,'purpose':'仅供生成器回归验证，不是登记项目'}
    manual={'chapters':[
        {'title':'一、软件概述','blocks':[{'type':'paragraph','text':'这是隔离的工具测试文字，图片是测试夹具，不作为真实运行或软著申请证据。'}]},
        {'title':'二、环境','blocks':[{'type':'table','headers':['项目','说明'],'rows':[['输入','测试用文字与源码'],['输出','工具验证产物']]}]},
        {'title':'三、业务流程','blocks':[{'type':'diagram','path':'fixture.png','caption':'生成器验证图片','evidence':'test fixture only'}]},
        {'title':'四、操作说明','blocks':[{'type':'paragraph','text':'生成器只验证排版、回读与哈希，不能判断软件业务真实性。'},
          {'type':'screenshot','path':'fixture.png','caption':'明确标识的测试图片','evidence':'test fixture only; not a real screenshot'}]}
    ]}
    input_path=base/'input.json';input_path.write_text(json.dumps({'schema':1,'application':app,'source_snapshot':'source.json','manual':manual},ensure_ascii=False),encoding='utf-8')
    return input_path,project

class MaterialsTests(unittest.TestCase):
    def test_unmodified_schema_example_is_not_a_project(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):build(ROOT/'assets/material-input.example.json',Path(tmp)/'out')
            self.assertFalse((Path(tmp)/'out').exists())

    def test_independent_generation_and_exact_roundtrip(self):
        with tempfile.TemporaryDirectory() as tmp:
            p,project=fixture(tmp);out=Path(tmp)/'out'
            with patch.object(Path,'home',side_effect=AssertionError('No other skill directory lookup')):
                result=build(p,out);report=audit(out)
            self.assertTrue(result['static_ok']);self.assertTrue(report['static_ok'])
            self.assertFalse(report['submission_ready']);self.assertGreater(result['missing_field_count'],0)
            self.assertIn('  significant spaces',(out/'代码材料_全部.txt').read_text(encoding='utf-8'))
            self.assertEqual((project/'main.py').read_text(encoding='utf-8'),(out/'代码材料_全部.txt').read_text(encoding='utf-8'))

    def test_word_tampering_detects_spaces(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            p,_=fixture(tmp);out=Path(tmp)/'out';build(p,out)
            doc=Document(out/'代码材料_全部.docx')
            target=next(x for x in doc.paragraphs if 'significant spaces' in x.text)
            target.text=target.text.replace('  significant',' significant');doc.save(out/'代码材料_全部.docx')
            report=audit(out)
            self.assertTrue(any(x['kind']=='source_roundtrip_mismatch' for x in report['issues']))

    def test_drift_blocks_generation(self):
        with tempfile.TemporaryDirectory() as tmp:
            p,project=fixture(tmp);(project/'main.py').write_text('changed',encoding='utf-8')
            with self.assertRaises(ValueError):build(p,Path(tmp)/'out')
            self.assertFalse((Path(tmp)/'out').exists())

    def test_existing_outputs_are_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            p,_=fixture(tmp);out=Path(tmp)/'out';build(p,out);before=(out/'操作手册.docx').read_bytes()
            with self.assertRaises(FileExistsError):build(p,out)
            self.assertEqual(before,(out/'操作手册.docx').read_bytes())

    def test_missing_screenshot_is_not_success(self):
        with tempfile.TemporaryDirectory() as tmp:
            p,_=fixture(tmp);(Path(tmp)/'fixture.png').unlink()
            with self.assertRaises(ValueError):build(p,Path(tmp)/'out')
            self.assertTrue((Path(tmp)/'out/生成未完成.txt').exists())

    def test_mapping_order_tampering_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            p,_=fixture(tmp);out=Path(tmp)/'out';build(p,out);m=read_json(out/'生成记录.json')
            m['source_mapping'][0]['line']=999
            (out/'生成记录.json').write_text(json.dumps(m),encoding='utf-8')
            self.assertFalse(audit(out)['static_ok'])

    def test_control_characters_are_not_silently_dropped(self):
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp);p,project=fixture(tmp)
            (project/'main.py').write_text('print("x")\x0b',encoding='utf-8')
            (base/'source.json').unlink();snapshot(project,base/'selection.json',base/'source.json')
            with self.assertRaises(ValueError):build(p,base/'out')
            self.assertFalse((base/'out').exists())

    def test_long_unicode_tabs_and_numbering_preserve_text(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            p,_=fixture(tmp,True);out=Path(tmp)/'out';r=build(p,out,40)
            self.assertTrue(audit(out)['static_ok']);self.assertGreater(r['logical_code_pages'],2)
            doc=Document(out/'代码材料_全部.docx')
            self.assertTrue(doc.paragraphs[50].text.startswith('01 '));self.assertTrue(doc.paragraphs[50].paragraph_format.page_break_before)
            value='\t中文  a\t'+('x '*100)
            self.assertEqual(''.join(wrap_line(value,30)),value)

    def test_manual_table_mutation_detected(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            p,_=fixture(tmp);out=Path(tmp)/'out';build(p,out)
            doc=Document(out/'操作手册.docx');doc.tables[0].cell(1,1).text='unrelated';doc.save(out/'操作手册.docx')
            self.assertTrue(any(x['kind']=='manual_text_missing_or_changed' for x in audit(out)['issues']))

    def test_header_identity_is_audited(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            p,_=fixture(tmp);out=Path(tmp)/'out';build(p,out)
            doc=Document(out/'代码材料_全部.docx');doc.sections[0].header.paragraphs[0].text='wrong name';doc.save(out/'代码材料_全部.docx')
            self.assertTrue(any(x['kind']=='header_identity_mismatch' for x in audit(out)['issues']))

    def test_continuous_pdf_ranges_and_source_unchanged(self):
        from pypdf import PdfWriter,PdfReader
        with tempfile.TemporaryDirectory() as tmp:
            base=Path(tmp);source=base/'full.pdf';writer=PdfWriter()
            for i in range(73):writer.add_blank_page(width=100+i,height=200)
            with source.open('wb') as s:writer.write(s)
            before=source.read_bytes();r=export(source,base/'deposit.pdf')
            self.assertEqual(r['source_pages_1based'],list(range(1,31))+list(range(44,74)))
            self.assertEqual([int(x.mediabox.width) for x in PdfReader(base/'deposit.pdf').pages],list(range(100,130))+list(range(143,173)))
            self.assertEqual(source.read_bytes(),before)
            with self.assertRaises(FileExistsError):export(source,base/'deposit.pdf')

    def test_project_scaffolding_is_empty_and_protected(self):
        with tempfile.TemporaryDirectory() as tmp:
            r=initialize(tmp,'new-project');self.assertFalse(r['software_implemented'])
            self.assertEqual(list((Path(r['project'])/'src').iterdir()),[])
            with self.assertRaises(FileExistsError):initialize(tmp,'new-project')
            for name in ['../escape','CON','trailing.']:
                with self.assertRaises(ValueError):initialize(tmp,name)

if __name__=='__main__':unittest.main()
